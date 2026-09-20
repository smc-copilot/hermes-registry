#!/usr/bin/env python3
"""
Convert a Markdown customer profile report to a professional Word document.
Usage: python3 md-to-docx-converter.py <input.md> [output.docx]
If output.docx is omitted, derives from input filename.
"""

import re
import sys
import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

def set_cell_shading(cell, color):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shading = OxmlElement('w:shd')
    shading.set(qn('w:fill'), color)
    shading.set(qn('w:val'), 'clear')
    tcPr.append(shading)

def add_styled_table(doc, headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.alignment = 1
    table.style = 'Table Grid'
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ''
        p = cell.paragraphs[0]
        run = p.add_run(h)
        run.bold = True
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_cell_shading(cell, '2F5496')
    for r_idx, row in enumerate(rows):
        for c_idx, val in enumerate(row):
            cell = table.rows[r_idx + 1].cells[c_idx]
            cell.text = ''
            p = cell.paragraphs[0]
            run = p.add_run(str(val))
            run.font.size = Pt(9)
            if r_idx % 2 == 1:
                set_cell_shading(cell, 'D6E4F0')
    doc.add_paragraph()
    return table

def parse_md_table(lines, start_idx):
    headers = [c.strip() for c in lines[start_idx].split('|')[1:-1]]
    rows = []
    i = start_idx + 2
    while i < len(lines) and lines[i].strip().startswith('|'):
        row = [c.strip() for c in lines[i].split('|')[1:-1]]
        rows.append(row)
        i += 1
    return headers, rows, i

def main():
    if len(sys.argv) < 2:
        print("Usage: python3 md-to-docx-converter.py <input.md> [output.docx]")
        sys.exit(1)

    src = sys.argv[1]
    if not os.path.exists(src):
        print(f"Error: {src} not found")
        sys.exit(1)

    if len(sys.argv) >= 3:
        dst = sys.argv[2]
    else:
        dst = os.path.splitext(src)[0] + '.docx'

    doc = Document()
    for section in doc.sections:
        section.top_margin = Cm(2.5)
        section.bottom_margin = Cm(2.5)
        section.left_margin = Cm(2.8)
        section.right_margin = Cm(2.8)

    style = doc.styles['Normal']
    font = style.font
    font.name = '宋体'
    font.size = Pt(11)
    style.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')

    with open(src, 'r', encoding='utf-8') as f:
        content = f.read()

    lines = content.split('\n')
    i = 0
    in_blockquote = False
    blockquote_lines = []
    in_code_block = False

    while i < len(lines):
        line = lines[i]

        if line.strip().startswith('```'):
            in_code_block = not in_code_block
            i += 1
            continue

        if in_code_block:
            i += 1
            continue

        if line.strip().startswith('>'):
            if not in_blockquote:
                in_blockquote = True
                blockquote_lines = []
            blockquote_lines.append(line.strip()[1:].strip())
            i += 1
            continue
        elif in_blockquote:
            text = ' '.join(blockquote_lines)
            p = doc.add_paragraph()
            run = p.add_run(text)
            run.font.size = Pt(10)
            run.italic = True
            run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
            in_blockquote = False
            blockquote_lines = []

        if '|' in line and line.strip().startswith('|') and i + 1 < len(lines) and '---' in lines[i + 1]:
            headers, rows, next_i = parse_md_table(lines, i)
            add_styled_table(doc, headers, rows)
            i = next_i
            continue

        if line.startswith('# '):
            h = doc.add_heading(line[2:].strip(), level=1)
            for run in h.runs:
                run.font.name = '黑体'
                run._element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')
            i += 1
            continue
        elif line.startswith('## '):
            h = doc.add_heading(line[3:].strip(), level=2)
            for run in h.runs:
                run.font.name = '黑体'
                run._element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')
            i += 1
            continue
        elif line.startswith('### '):
            h = doc.add_heading(line[4:].strip(), level=3)
            for run in h.runs:
                run.font.name = '黑体'
                run._element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')
            i += 1
            continue

        if line.strip() == '---':
            doc.add_paragraph('─' * 60)
            i += 1
            continue

        if not line.strip():
            i += 1
            continue

        text = line.strip()
        p = doc.add_paragraph()
        parts = re.split(r'(\*\*.*?\*\*)', text)
        for part in parts:
            if part.startswith('**') and part.endswith('**'):
                run = p.add_run(part[2:-2])
                run.bold = True
            else:
                run = p.add_run(part)
            run.font.size = Pt(11)
            run.font.name = '宋体'
            run._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')

        pf = p.paragraph_format
        pf.space_after = Pt(6)
        pf.line_spacing = 1.5
        i += 1

    doc.add_paragraph()
    p = doc.add_paragraph()
    run = p.add_run('免责声明：本报告由AI生成，数据来源于公开渠道，仅供内部参考，不构成投资或合作决策依据。')
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0x99, 0x99, 0x99)
    run.italic = True

    doc.save(dst)
    char_count = sum(len(p.text) for p in doc.paragraphs)
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                char_count += len(cell.text)
    print(f'✅ DOCX saved: {dst}  ({char_count} chars)')

if __name__ == '__main__':
    main()
