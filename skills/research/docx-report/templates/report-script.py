#!/usr/bin/env python3
"""
Word文档(.docx)生成基础模板 — docx-report skill
基于python-docx，支持中文字体、深蓝表头、交替行底色、封面页等专业格式。

使用方式：
1. 复制此脚本
2. 按需修改内容（替换所有 {{PLACEHOLDER}} 标记）
3. 运行: python scripts/xxx_report.py
"""

from docx import Document
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import datetime

doc = Document()

# ============ 样式设置 ============
style = doc.styles['Normal']
font = style.font
font.name = '宋体'
font.size = Pt(11)
style.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')

# 页边距
for section in doc.sections:
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(2.8)
    section.right_margin = Cm(2.8)

# ============ 辅助函数 ============
def set_cell_shading(cell, color):
    """设置单元格底色"""
    shading_elm = OxmlElement('w:shd')
    shading_elm.set(qn('w:fill'), color)
    cell._tc.get_or_add_tcPr().append(shading_elm)

def add_heading_styled(text, level=1):
    """添加黑体标题"""
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.name = '黑体'
        run.element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')
    return h

def add_para(text, bold=False, indent=True, color=None):
    """添加宋体正文段落"""
    p = doc.add_paragraph()
    if indent:
        p.paragraph_format.first_line_indent = Pt(22)
    p.paragraph_format.line_spacing = 1.5
    run = p.add_run(text)
    run.font.name = '宋体'
    run.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    run.font.size = Pt(11)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)
    return p

def add_table(headers, rows):
    """添加专业样式表格（深蓝表头+交替行底色）"""
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    # 表头
    HEADER_BG = '1E2761'  # 深蓝
    ROW_ALT_BG = 'E8EDF5'  # 浅蓝灰
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = h
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in p.runs:
                run.bold = True
                run.font.size = Pt(9)
                run.font.color.rgb = RGBColor(255, 255, 255)
        set_cell_shading(cell, HEADER_BG)

    # 数据行
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = table.rows[ri + 1].cells[ci]
            cell.text = str(val)
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(9)
            if ri % 2 == 1:
                set_cell_shading(cell, ROW_ALT_BG)

    doc.add_paragraph()  # 表后间距
    return table

def add_page_break():
    doc.add_page_break()


# ============================================================
# 封面（根据实际报告修改标题和副标题）
# ============================================================
for _ in range(6):
    doc.add_paragraph()

title_p = doc.add_paragraph()
title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = title_p.add_run('{{报告标题}}')
run.font.size = Pt(28)
run.bold = True
run.font.name = '黑体'
run.element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')
run.font.color.rgb = RGBColor(0x1E, 0x27, 0x61)

sub_p = doc.add_paragraph()
sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = sub_p.add_run('{{副标题/股票代码}}')
run.font.size = Pt(20)
run.font.name = '黑体'
run.element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')
run.font.color.rgb = RGBColor(0x1E, 0x27, 0x61)

doc.add_paragraph()

desc_p = doc.add_paragraph()
desc_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = desc_p.add_run('{{报告描述}}')
run.font.size = Pt(14)

doc.add_paragraph()
doc.add_paragraph()

date_p = doc.add_paragraph()
date_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
today_str = datetime.date.today().strftime("%Y年%m月%d日")
run = date_p.add_run(f'报告日期：{today_str}')
run.font.size = Pt(12)

disc_p = doc.add_paragraph()
disc_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = disc_p.add_run('{{数据来源声明}}')
run.font.size = Pt(10)
run.font.color.rgb = RGBColor(128, 128, 128)

disc2_p = doc.add_paragraph()
disc2_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = disc2_p.add_run('免责声明：本报告所有分析仅供参考，不构成任何投资建议')
run.font.size = Pt(10)
run.font.color.rgb = RGBColor(128, 128, 128)

doc.add_page_break()

# ============================================================
# 目录
# ============================================================
add_heading_styled('目  录', level=1)
# {{在此列出各章节标题}}

doc.add_page_break()

# ============================================================
# 正文章节（按需复制以下模板）
# ============================================================
add_heading_styled('{{章节标题}}', level=1)
add_heading_styled('{{小节标题}}', level=2)
add_para('{{正文内容，支持多段落。财务数据需标注来源和置信度。推算数据标注[推导估算值]。}}')

# 表格示例
# add_table(
#     ['列1', '列2', '列3'],
#     [
#         ['数据1', '数据2', '数据3'],
#         ['数据4', '数据5', '数据6'],
#     ]
# )

# ============================================================
# 附录
# ============================================================
add_page_break()
add_heading_styled('附录：主要数据来源', level=1)
# add_table(...)

add_para('')
add_para('免责声明：本报告基于公开信息撰写，所有分析和结论仅供参考，不构成任何投资建议。投资者应当独立判断并自行承担投资风险。')
add_para(f'报告生成时间：{datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")}')

# ============ 保存 ============
output_path = '{{输出路径}}'
doc.save(output_path)
print(f'报告已保存至: {output_path}')
